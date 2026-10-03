import os
import re
import json
import psycopg2
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from typing import Dict, List, Any, Optional

def get_pg_connection(cfg: dict):
    return psycopg2.connect(
        host=cfg.get("postgres_host", "192.168.10.237"),
        port=int(cfg.get("postgres_port", 5432)),
        dbname=cfg.get("postgres_db", "tally_reconciliation_db"),
        user=cfg.get("postgres_user", "postgres"),
        password=cfg.get("postgres_pass", "Laduma@2026"),
        connect_timeout=6
    )

def clean_num(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except:
        c = re.sub(r'[^0-9.-]', '', str(v).strip())
        try:
            return float(c)
        except:
            return 0.0

def clean_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()

class RBMExcelReader:
    def __init__(self, filepath: str):
        self.wb = openpyxl.load_workbook(filepath, data_only=True)

    def get_trading_data(self) -> List[Dict[str, Any]]:
        sh_names = [s for s in self.wb.sheetnames if "TRADING" in s.upper()]
        if not sh_names:
            return []
        ws = self.wb[sh_names[0]]
        records = []
        for r in range(3, ws.max_row + 1):
            branch = clean_text(ws.cell(r, 1).value)
            if not branch or branch.upper() in ["TOTAL", "GRAND TOTAL"]:
                continue
            records.append({
                "branch": branch,
                "opening_stock_val": clean_num(ws.cell(r, 3).value),
                "purchase_val": clean_num(ws.cell(r, 5).value),
                "purchase_rtn_val": clean_num(ws.cell(r, 7).value),
                "sales_val": clean_num(ws.cell(r, 9).value),
                "sales_cost": clean_num(ws.cell(r, 10).value),
                "sales_rtn_val": clean_num(ws.cell(r, 12).value),
                "sales_rtn_cost": clean_num(ws.cell(r, 13).value),
                "ibt_in_val": clean_num(ws.cell(r, 15).value),
                "ibt_out_val": clean_num(ws.cell(r, 17).value),
                "ibt_reject_receive_val": clean_num(ws.cell(r, 19).value),
                "stock_take_in_val": clean_num(ws.cell(r, 21).value),
                "stock_take_out_val": clean_num(ws.cell(r, 23).value),
                "stock_consumption_val": clean_num(ws.cell(r, 25).value),
                "net_sales_cost": clean_num(ws.cell(r, 10).value) - clean_num(ws.cell(r, 13).value),
                "net_purchase_val": clean_num(ws.cell(r, 5).value) - clean_num(ws.cell(r, 7).value)
            })
        return records

    def get_branch_reports_data(self) -> List[Dict[str, Any]]:
        target_name = None
        for name in ["TALLY DETAILSHEET", "COS-SOFTWARE DETAILSHEET"]:
            if name in self.wb.sheetnames:
                target_name = name
                break
        if not target_name:
            return []
        ws = self.wb[target_name]
        start_row = 3
        for r in range(1, 10):
            val = clean_text(ws.cell(r, 1).value).upper()
            if "BRANCH REPORTS" in val or "OTHER REPORTS" in val:
                start_row = r + 1
                break
        records = []
        for r in range(start_row, ws.max_row + 1):
            bname = clean_text(ws.cell(r, 1).value)
            if not bname or any(k in bname.upper() for k in ["TOTAL", "TRADING ACCOUNT", "TALLY"]):
                break
            records.append({
                "branch": bname,
                "opening_stock": clean_num(ws.cell(r, 2).value),
                "total_purchase": clean_num(ws.cell(r, 3).value),
                "ibt_in": clean_num(ws.cell(r, 4).value),
                "ibt_out": clean_num(ws.cell(r, 5).value),
                "ibt_out_reject_prev": clean_num(ws.cell(r, 6).value),
                "ibt_out_rejected": clean_num(ws.cell(r, 7).value),
                "ibt_out_reject_next": clean_num(ws.cell(r, 8).value),
                "stock_discrepancy": clean_num(ws.cell(r, 9).value),
                "stock_consumption": clean_num(ws.cell(r, 10).value),
                "damage_con": clean_num(ws.cell(r, 11).value),
                "merge_split": clean_num(ws.cell(r, 12).value),
                "stock_in_out": clean_num(ws.cell(r, 13).value),
                "manufacturing": clean_num(ws.cell(r, 14).value),
                "closing_stock": clean_num(ws.cell(r, 15).value),
                "cos_as_per_reports": clean_num(ws.cell(r, 16).value),
                "cos_in_sales_report": clean_num(ws.cell(r, 17).value),
                "diff": clean_num(ws.cell(r, 18).value)
            })
        return records

def calculate_reconciliation(branch_name: str, tr: dict, br: dict, tal: dict) -> dict:
    br_opening = br.get("opening_stock", 0.0)
    br_purchase = br.get("total_purchase", 0.0)
    br_ibt_in = br.get("ibt_in", 0.0)
    br_ibt_out = br.get("ibt_out", 0.0)
    br_disc = br.get("stock_discrepancy", 0.0)
    br_consump = br.get("stock_consumption", 0.0)
    br_damage = br.get("damage_con", 0.0)
    br_merge = br.get("merge_split", 0.0)
    br_stock_in_out = br.get("stock_in_out", 0.0)
    br_mfg = br.get("manufacturing", 0.0)
    br_closing = br.get("closing_stock", 0.0)
    br_reject_prev = br.get("ibt_out_reject_prev", 0.0)

    cos_branch_reports = (
        br_opening + br_purchase + br_ibt_in - br_ibt_out + br_disc -
        br_consump + br_damage + br_merge + br_stock_in_out + br_mfg -
        br_closing + br_reject_prev
    )

    cos_sales_report = (
        tr.get("sales_cost", 0.0) - tr.get("sales_rtn_cost", 0.0)
        if tr else br.get("cos_in_sales_report", 0.0)
    )

    tr_opening = tr.get("opening_stock_val", br_opening)
    tr_purchase = tr.get("net_purchase_val", br_purchase)
    tr_ibt_in = tr.get("ibt_in_val", br_ibt_in)
    tr_ibt_out = tr.get("ibt_out_val", br_ibt_out)
    tr_consump = tr.get("stock_consumption_val", br_consump)

    cos_trading_account = (
        tr_opening + tr_purchase + tr_ibt_in - tr_ibt_out + br_disc -
        tr_consump + br_damage + br_merge + br_stock_in_out + br_mfg - br_closing
    )

    t_opening = tal.get("opening_stock", br_opening)
    t_purchase = tal.get("purchases", 0.0)
    t_ibt_in = tal.get("ibt_in", 0.0)
    t_ibt_out = tal.get("ibt_out", 0.0)
    t_consump = tal.get("stock_consumed", 0.0) + tal.get("ho_stock_consumed", 0.0)
    t_closing = tal.get("closing_stock", br_closing)

    cos_tally = (
        t_opening + t_purchase + t_ibt_in - t_ibt_out + br_disc -
        t_consump + br_damage + br_merge + br_stock_in_out + br_mfg - t_closing
    )

    diff_reports_vs_sales = cos_branch_reports - cos_sales_report
    diff_trading_vs_tally = cos_trading_account - cos_tally
    diff_reports_vs_tally = cos_branch_reports - cos_tally

    is_flagged = abs(diff_reports_vs_sales) > 1.0 or abs(diff_reports_vs_tally) > 1.0
    status = "MATCHED" if not is_flagged else "MISMATCH"

    return {
        "branch_name": branch_name,
        "cos_branch_reports": round(cos_branch_reports, 2),
        "cos_trading_account": round(cos_trading_account, 2),
        "cos_sales_report": round(cos_sales_report, 2),
        "cos_tally": round(cos_tally, 2),
        "diff_reports_vs_sales": round(diff_reports_vs_sales, 2),
        "diff_trading_vs_tally": round(diff_trading_vs_tally, 2),
        "diff_reports_vs_tally": round(diff_reports_vs_tally, 2),
        "is_flagged": is_flagged,
        "status": status,
        "details": {
            "branch_reports": br,
            "trading_account": tr,
            "tally": tal
        }
    }

def generate_reconciled_excel(rows: List[dict], output_path: str, period_label: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COS Discrepancy Summary"

    hdr_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    flag_fill = PatternFill(start_color="FFD8D8", end_color="FFD8D8", fill_type="solid")
    match_fill = PatternFill(start_color="D8F3D8", end_color="D8F3D8", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3')
    )

    headers = [
        "Branch Name", "COS (Branch Reports)", "COS (Trading Account)",
        "COS (Sales Cost)", "COS (Tally)", "Diff (Reports vs Sales)",
        "Diff (Reports vs Tally)", "Audit Status"
    ]
    ws.append(headers)

    for col_idx, col_name in enumerate(headers, 1):
        c = ws.cell(1, col_idx)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        row_data = [
            r["branch_name"], r["cos_branch_reports"], r["cos_trading_account"],
            r["cos_sales_report"], r["cos_tally"], r["diff_reports_vs_sales"],
            r["diff_reports_vs_tally"], r["status"]
        ]
        ws.append(row_data)
        cur_row = ws.max_row
        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(cur_row, c_idx)
            cell.border = border_thin
            if c_idx == 8:
                cell.fill = flag_fill if r["is_flagged"] else match_fill
                cell.alignment = Alignment(horizontal="center")
            elif c_idx > 1:
                cell.number_format = '#,##0.00'

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(output_path)


# ── TALLY 6 TARGET GROUPS RECONCILIATION ENGINE ────────────────────────────

TALLY_TARGET_GROUPS = [
    "Purchase Accounts",
    "Head Office Stock Consumed by Branch",
    "Stock Consumed By Branch",
    "Stock Transfers Branches",
    "Stock Discrepancy in Branches",
    "Damage Conversion Accounts"
]

def get_ledger_branch_mappings(cfg: dict) -> Dict[str, dict]:
    mappings = {}
    try:
        conn = get_pg_connection(cfg)
        cur = conn.cursor()
        cur.execute("SELECT group_name, ledger_name, branch_name, is_verified FROM tally_ledger_mappings;")
        for g, l, b, is_ver in cur.fetchall():
            info = {
                "branch": (b or "").strip(),
                "is_verified": bool(is_ver)
            }
            key = f"{g.strip().lower()}::{l.strip().lower()}"
            mappings[key] = info
            mappings[l.strip().lower()] = info
        conn.close()
    except Exception as e:
        print(f"Notice: unable to load tally_ledger_mappings: {e}")
    return mappings

def smart_match_branch(group: str, ledger: str) -> str:
    l = ledger.lower().strip()
    if 'boards wr' in l or 'board division' in l or re.search(r'\bboards\b', l): return 'Boards WR'
    if 'headoffice' in l or 'head office' in l or re.search(r'\bho\b', l): return 'HeadOffice WR'
    if 'warehouse-bf' in l or 'warehouse bf' in l or 'warehouse - bf' in l: return 'Warehouse BF'
    if 'warehouse wr' in l or 'warehouse-wr' in l: return 'Warehouse WR'
    if 'distribution centre' in l or 'dc import' in l: return 'Distribution Centre'
    if 'msholozi' in l: return 'Msholozi'
    if 'rockysdrift' in l or 'rockys drift' in l: return 'Rockys Drift'
    if 'lebo warehouse' in l: return 'Lebo Warehouse'
    if 'lebo' in l or 'lebowakgomo' in l: return 'Lebowakgomo'
    if 'mega store' in l or 'megastore' in l or re.search(r'\bmega\b', l): return 'Megastore'
    if 'violet bank' in l or 'violetbank' in l: return 'Violetbank'
    if 'chuenespoort' in l or 'chuenspoort' in l or re.search(r'\bchuenes\b', l): return 'Chuenespoort'
    if 'malamulale' in l or 'malamulele' in l or re.search(r'\bmala\b', l): return 'Malamulele'
    if 'thulamahashe' in l or re.search(r'\bthul\b', l): return 'Thulamahashe'
    if 'acornhoek' in l or re.search(r'\bacorn\b', l): return 'Acornhoek'
    if 'bushbuckridge' in l or re.search(r'\bbush\b', l): return 'Bushbuckridge'
    if 'burgersfort' in l: return 'Burgersfort'
    if 'driekop' in l: return 'Driekop'
    if 'elukwatini' in l or 'elukwathini' in l: return 'Elukwatini'
    if 'giyani' in l: return 'Giyani'
    if 'hluvukani' in l or re.search(r'\bhluv\b', l): return 'Hluvukani'
    if 'janefurse' in l or 'jane furse' in l: return 'Janefurse'
    if 'mkhuhlu' in l or re.search(r'\bmkhu\b', l): return 'Mkhuhlu'
    if 'moratiwa' in l: return 'Moratiwa'
    if 'nebo' in l: return 'Nebo'
    if 'nobody' in l: return 'Nobody'
    if 'phola park' in l or 'phola' in l: return 'Phola Park'
    if 'seshego' in l or re.search(r'\bsesh\b', l): return 'Seshego'
    if 'swalala' in l: return 'Swalala'
    if 'tonga' in l: return 'Tonga'
    if 'bochum' in l: return 'Bochum'
    if 'builden' in l: return 'Builden'
    if 'dennilton' in l: return 'Dennilton'

    # Fallback to general regex extraction
    n = ledger.strip()
    n = re.sub(r'^(HO\s+Stock\s+Consumed\s+by\s+Branch\s*[-–]\s*)', '', n, flags=re.IGNORECASE)
    n = re.sub(r'^(Stock\s+Consumed\s+by\s+(?:Branch\s*)?[-–]\s*)', '', n, flags=re.IGNORECASE)
    n = re.sub(r'^(Stock\s+discrepancy\s*(?:in\s+Branches)?\s*[-–]?\s*)', '', n, flags=re.IGNORECASE)
    n = re.sub(r'^(Damage\s+Conversion\s*(?:Accounts|Account)?\s*[-–]?\s*)', '', n, flags=re.IGNORECASE)
    n = re.sub(r'^(IBT\s+)', '', n, flags=re.IGNORECASE)
    n = re.sub(r'^(Purchase\s*[-–]?\s*)', '', n, flags=re.IGNORECASE)
    return n.strip()

def extract_branch_from_ledger(name: str) -> str:
    return smart_match_branch("", name)

def pull_tally_group_summaries(tally_url: str, company: str, from_date_str: str, to_date_str: str, cfg: Optional[dict] = None) -> dict:
    import xml.etree.ElementTree as ET
    import requests

    db_mappings = {}
    if cfg:
        db_mappings = get_ledger_branch_mappings(cfg)

    f_date = from_date_str.replace("-", "").replace("/", "")[:8]
    t_date = to_date_str.replace("-", "").replace("/", "")[:8]
    clean_comp = company.replace("&", "&amp;")

    all_items = []
    groups_summary = {}

    for grp in TALLY_TARGET_GROUPS:
        targets = [grp]
        if grp == "Purchase Accounts":
            req_sub = f"""<ENVELOPE>
              <HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST><TYPE>Collection</TYPE><ID>SubGrps</ID></HEADER>
              <BODY>
                <DESC>
                  <STATICVARIABLES>
                    <SVCURRENTCOMPANY>{clean_comp}</SVCURRENTCOMPANY>
                    <SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>
                  </STATICVARIABLES>
                  <TDL>
                    <TDLMESSAGE>
                      <COLLECTION NAME="SubGrps">
                        <TYPE>Group</TYPE>
                        <CHILDOF>Purchase Accounts</CHILDOF>
                        <FETCH>NAME</FETCH>
                      </COLLECTION>
                    </TDLMESSAGE>
                  </TDL>
                </DESC>
              </BODY>
            </ENVELOPE>"""
            try:
                r_sub = requests.post(tally_url, data=req_sub.encode('utf-8'), timeout=15)
                if r_sub.status_code == 200:
                    root_sub = ET.fromstring(r_sub.text)
                    for g in root_sub.findall(".//GROUP"):
                        gname = g.findtext("NAME") or g.get("NAME")
                        if not gname:
                            continue
                        gn = gname.strip()
                        if gn.upper().startswith("IBT") or "INTER BRANCH" in gn.upper():
                            continue
                        targets.append(gn)
            except Exception as e:
                print(f"Error fetching sub-groups for {grp}: {e}")

        grp_debit = 0.0
        grp_credit = 0.0
        grp_closing = 0.0
        grp_opening = 0.0
        seen_ledgers = set()

        try:
            for target in targets:
                req_xml = f"""<ENVELOPE>
                  <HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST><TYPE>Collection</TYPE><ID>GroupLedgers</ID></HEADER>
                  <BODY>
                    <DESC>
                      <STATICVARIABLES>
                        <SVCURRENTCOMPANY>{clean_comp}</SVCURRENTCOMPANY>
                        <SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>
                        <SVFROMDATE TYPE="Date">{f_date}</SVFROMDATE>
                        <SVTODATE TYPE="Date">{t_date}</SVTODATE>
                      </STATICVARIABLES>
                      <TDL>
                        <TDLMESSAGE>
                          <COLLECTION NAME="GroupLedgers">
                            <TYPE>Ledger</TYPE>
                            <CHILDOF>{target}</CHILDOF>
                            <FETCH>NAME, OPENINGBALANCE, DEBITTOTALS, CREDITTOTALS, CLOSINGBALANCE</FETCH>
                          </COLLECTION>
                        </TDLMESSAGE>
                      </TDL>
                    </DESC>
                  </BODY>
                </ENVELOPE>"""

                r = requests.post(tally_url, data=req_xml.encode('utf-8'), headers={'Content-Type': 'text/xml'}, timeout=25)
                if r.status_code != 200:
                    continue
                root = ET.fromstring(r.text)
                ledgers = root.findall(".//LEDGER")

                for l in ledgers:
                    name = l.get("NAME") or l.findtext("NAME")
                    if not name or not name.strip():
                        continue
                    name = name.strip()
                    if name.lower() in seen_ledgers:
                        continue
                    seen_ledgers.add(name.lower())

                    def parse_f(txt):
                        if not txt: return 0.0
                        try: return float(txt.strip())
                        except: return 0.0

                    op = parse_f(l.findtext("OPENINGBALANCE"))
                    dr_raw = parse_f(l.findtext("DEBITTOTALS"))
                    cr_raw = parse_f(l.findtext("CREDITTOTALS"))
                    cl = parse_f(l.findtext("CLOSINGBALANCE"))

                    dr = abs(dr_raw)
                    cr = abs(cr_raw)
                    
                    m_key = f"{grp.strip().lower()}::{name.lower()}"
                    is_mapped = False
                    is_verified = False
                    if m_key in db_mappings and db_mappings[m_key].get("branch"):
                        branch = db_mappings[m_key]["branch"]
                        is_verified = db_mappings[m_key].get("is_verified", False)
                        is_mapped = True
                    elif name.lower() in db_mappings and db_mappings[name.lower()].get("branch"):
                        branch = db_mappings[name.lower()]["branch"]
                        is_verified = db_mappings[name.lower()].get("is_verified", False)
                        is_mapped = True
                    else:
                        branch = smart_match_branch(grp, name)
                        is_verified = False
                        is_mapped = bool(branch and branch != "UNMAPPED")

                    item = {
                        "group_name": grp,
                        "ledger_name": name,
                        "branch_name": branch if branch != "UNMAPPED" else "",
                        "is_mapped": is_mapped,
                        "is_verified": is_verified,
                        "is_new": not is_verified,
                        "opening_balance": round(op, 2),
                        "debit_amount": round(dr, 2),
                        "credit_amount": round(cr, 2),
                        "closing_balance": round(cl, 2)
                    }
                    all_items.append(item)

                    grp_opening += op
                    grp_debit += dr
                    grp_credit += cr
                    grp_closing += cl

            groups_summary[grp] = {
                "group_name": grp,
                "total_opening": round(grp_opening, 2),
                "total_debit": round(grp_debit, 2),
                "total_credit": round(grp_credit, 2),
                "total_closing": round(grp_closing, 2),
                "ledger_count": len(seen_ledgers)
            }
        except Exception as e:
            print(f"Error pulling Tally group '{grp}': {e}")
            groups_summary[grp] = {
                "group_name": grp,
                "total_opening": 0.0,
                "total_debit": 0.0,
                "total_credit": 0.0,
                "total_closing": 0.0,
                "ledger_count": 0,
                "error": str(e)
            }

    return {
        "items": all_items,
        "groups_summary": groups_summary,
        "total_ledgers": len(all_items)
    }

def save_tally_group_summaries(cfg: dict, period_label: str, from_date: str, to_date: str, company: str, items: List[dict]):
    conn = get_pg_connection(cfg)
    cur = conn.cursor()

    # Ensure period exists
    cur.execute("""
        INSERT INTO reconciliation_periods (period_label, from_date, to_date, company_name)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (period_label) DO UPDATE 
        SET from_date=EXCLUDED.from_date, to_date=EXCLUDED.to_date, updated_at=CURRENT_TIMESTAMP
        RETURNING id;
    """, (period_label, from_date, to_date, company))
    period_id = cur.fetchone()[0]

    # Save to tally_group_summaries
    for it in items:
        cur.execute("""
            INSERT INTO tally_group_summaries (
                period_id, period_label, from_date, to_date, group_name, ledger_name,
                opening_balance, debit_amount, credit_amount, closing_balance, updated_at
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
            ON CONFLICT (period_label, group_name, ledger_name) DO UPDATE
            SET opening_balance = EXCLUDED.opening_balance,
                debit_amount = EXCLUDED.debit_amount,
                credit_amount = EXCLUDED.credit_amount,
                closing_balance = EXCLUDED.closing_balance,
                updated_at = CURRENT_TIMESTAMP;
        """, (
            period_id, period_label, from_date, to_date,
            it["group_name"], it["ledger_name"],
            it["opening_balance"], it["debit_amount"], it["credit_amount"], it["closing_balance"]
        ))

    # Ensure any new ledger discovered from Tally is registered in tally_ledger_mappings as is_verified=FALSE
    cur.execute("SELECT group_name, ledger_name FROM tally_ledger_mappings;")
    existing_map = set((r[0].strip().lower(), r[1].strip().lower()) for r in cur.fetchall())

    for it in items:
        grp = it["group_name"]
        lname = it["ledger_name"]
        key = (grp.strip().lower(), lname.strip().lower())
        if key not in existing_map:
            b_sug = it.get("branch_name", "")
            if b_sug == "UNMAPPED":
                b_sug = ""
            cur.execute("""
                INSERT INTO tally_ledger_mappings (company_name, group_name, ledger_name, branch_name, is_verified, updated_at)
                VALUES (%s, %s, %s, %s, FALSE, CURRENT_TIMESTAMP)
                ON CONFLICT (group_name, ledger_name) DO NOTHING;
            """, (company, grp, lname, b_sug))
            existing_map.add(key)

    # Aggregate per branch and upsert into tally_extracted_data
    branch_map = {}
    for it in items:
        b = it.get("branch_name")
        if not b:
            continue
        if b not in branch_map:
            branch_map[b] = {
                "purchases": 0.0,
                "ibt_in": 0.0,
                "ibt_out": 0.0,
                "stock_consumed": 0.0,
                "ho_stock_consumed": 0.0,
                "damage_conversion": 0.0,
                "stock_discrepancy": 0.0,
                "opening_stock": 0.0,
                "closing_stock": 0.0,
                "details": []
            }
        grp = it["group_name"]
        branch_map[b]["details"].append(it)
        if grp == "Purchase Accounts":
            net_p = it["debit_amount"] - it["credit_amount"]
            if net_p == 0 and it.get("closing_balance"):
                net_p = abs(it["closing_balance"])
            branch_map[b]["purchases"] += net_p
        elif grp == "Stock Transfers Branches":
            branch_map[b]["ibt_in"] += it["debit_amount"]
            branch_map[b]["ibt_out"] += it["credit_amount"]
        elif grp == "Stock Consumed By Branch":
            branch_map[b]["stock_consumed"] += it["closing_balance"]
        elif grp == "Head Office Stock Consumed by Branch":
            branch_map[b]["ho_stock_consumed"] += it["closing_balance"]
        elif grp == "Damage Conversion Accounts":
            branch_map[b]["damage_conversion"] += it["closing_balance"]
        elif grp == "Stock Discrepancy in Branches":
            branch_map[b]["stock_discrepancy"] += it["closing_balance"]

    for bname, bdata in branch_map.items():
        cur.execute("""
            INSERT INTO tally_extracted_data (
                period_id, branch_name, purchases, ibt_in, ibt_out,
                stock_consumed, ho_stock_consumed, opening_stock, closing_stock,
                raw_details, updated_at
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
            ON CONFLICT (period_id, branch_name) DO UPDATE
            SET purchases = EXCLUDED.purchases,
                ibt_in = EXCLUDED.ibt_in,
                ibt_out = EXCLUDED.ibt_out,
                stock_consumed = EXCLUDED.stock_consumed,
                ho_stock_consumed = EXCLUDED.ho_stock_consumed,
                raw_details = EXCLUDED.raw_details,
                updated_at = CURRENT_TIMESTAMP;
        """, (
            period_id, bname, bdata["purchases"], bdata["ibt_in"], bdata["ibt_out"],
            bdata["stock_consumed"], bdata["ho_stock_consumed"],
            bdata["opening_stock"], bdata["closing_stock"],
            json.dumps(bdata["details"])
        ))

    conn.commit()
    conn.close()
    return period_id

def get_tally_group_summaries_from_db(cfg: dict, period_label: str) -> dict:
    conn = get_pg_connection(cfg)
    cur = conn.cursor()

    cur.execute("""
        SELECT group_name, ledger_name, opening_balance, debit_amount, credit_amount, closing_balance, updated_at, from_date, to_date
        FROM tally_group_summaries
        WHERE period_label = %s
        ORDER BY group_name, ledger_name;
    """, (period_label,))
    rows = cur.fetchall()

    if not rows:
        conn.close()
        return {
            "has_data": False,
            "period": period_label,
            "items": [],
            "groups_summary": {},
            "total_ledgers": 0,
            "updated_at": None
        }

    items = []
    groups_summary = {g: {
        "group_name": g,
        "total_opening": 0.0,
        "total_debit": 0.0,
        "total_credit": 0.0,
        "total_closing": 0.0,
        "ledger_count": 0
    } for g in TALLY_TARGET_GROUPS}

    last_ts = None
    f_date = None
    t_date = None

    db_mappings = get_ledger_branch_mappings(cfg)

    for r in rows:
        grp = r[0]
        lname = r[1]
        op = float(r[2] or 0)
        dr = float(r[3] or 0)
        cr = float(r[4] or 0)
        cl = float(r[5] or 0)
        ts = str(r[6]) if r[6] else None
        if not last_ts and ts:
            last_ts = ts
        if not f_date and r[7]:
            f_date = str(r[7])
        if not t_date and r[8]:
            t_date = str(r[8])

        m_key = f"{grp.strip().lower()}::{lname.lower()}"
        is_mapped = False
        is_verified = False
        if m_key in db_mappings and db_mappings[m_key].get("branch"):
            branch = db_mappings[m_key]["branch"]
            is_verified = db_mappings[m_key].get("is_verified", False)
            is_mapped = True
        elif lname.lower() in db_mappings and db_mappings[lname.lower()].get("branch"):
            branch = db_mappings[lname.lower()]["branch"]
            is_verified = db_mappings[lname.lower()].get("is_verified", False)
            is_mapped = True
        else:
            branch = smart_match_branch(grp, lname)
            is_verified = False
            is_mapped = bool(branch and branch != "UNMAPPED")

        item = {
            "group_name": grp,
            "ledger_name": lname,
            "branch_name": branch if branch != "UNMAPPED" else "",
            "is_mapped": is_mapped,
            "is_verified": is_verified,
            "is_new": not is_verified,
            "opening_balance": round(op, 2),
            "debit_amount": round(dr, 2),
            "credit_amount": round(cr, 2),
            "closing_balance": round(cl, 2)
        }
        items.append(item)

        if grp in groups_summary:
            groups_summary[grp]["total_opening"] += op
            groups_summary[grp]["total_debit"] += dr
            groups_summary[grp]["total_credit"] += cr
            groups_summary[grp]["total_closing"] += cl
            groups_summary[grp]["ledger_count"] += 1

    for g in groups_summary:
        groups_summary[g]["total_opening"] = round(groups_summary[g]["total_opening"], 2)
        groups_summary[g]["total_debit"] = round(groups_summary[g]["total_debit"], 2)
        groups_summary[g]["total_credit"] = round(groups_summary[g]["total_credit"], 2)
        groups_summary[g]["total_closing"] = round(groups_summary[g]["total_closing"], 2)

    conn.close()
    return {
        "has_data": True,
        "period": period_label,
        "from_date": f_date,
        "to_date": t_date,
        "items": items,
        "groups_summary": groups_summary,
        "total_ledgers": len(items),
        "updated_at": last_ts
    }

